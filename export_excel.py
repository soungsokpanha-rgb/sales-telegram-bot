from datetime import datetime
from pathlib import Path
from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import EXPORTS_DIR, CURRENCY_SYMBOL
import database as db


def generate_sales_excel(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None
) -> Path:
    """
    Generate professional Excel spreadsheet of daily sales records
    and summary across WWEG, WW271, and WWAE.
    """
    wb = openpyxl.Workbook()

    # 1. Sheet 1: Detailed Daily Records
    ws_records = wb.active
    ws_records.title = "Daily Sales Details"
    ws_records.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    title_font = Font(name="Calibri", size=15, bold=True, color="1F497D")
    sub_font = Font(name="Calibri", size=10, italic=True, color="595959")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    total_border = Border(
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000")
    )

    # Title Banner
    ws_records["A1"] = "CONCEPT STORES - DAILY SALES MASTER REPORT"
    ws_records["A1"].font = title_font
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws_records["A2"] = f"Generated on: {now_str} | Concept Stores: WWEG, WW271, WWAE"
    ws_records["A2"].font = sub_font

    headers = [
        "Date", "Store Code", "Store Name",
        "Actual Sales", "Target Sales", "Variance", "Target Hit %",
        "Transactions", "ATV (Avg Ticket)", "Footfall", "Submitted By", "Notes"
    ]

    header_row = 4
    for col_idx, h in enumerate(headers, start=1):
        cell = ws_records.cell(row=header_row, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    records = db.get_sales_records(start_date=start_date, end_date=end_date, limit=1000)

    current_row = header_row + 1
    for r in records:
        actual = r.get("actual_sales", 0.0)
        target = r.get("target_sales", 0.0)
        variance = actual - target
        pct = (actual / target) if target > 0 else 0.0
        txns = r.get("transactions", 0)
        atv = (actual / txns) if txns > 0 else 0.0

        ws_records.cell(row=current_row, column=1, value=r.get("report_date")).alignment = Alignment(horizontal="center")
        ws_records.cell(row=current_row, column=2, value=r.get("store_code")).alignment = Alignment(horizontal="center")
        ws_records.cell(row=current_row, column=3, value=r.get("store_name", ""))

        c_actual = ws_records.cell(row=current_row, column=4, value=actual)
        c_actual.number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'

        c_target = ws_records.cell(row=current_row, column=5, value=target)
        c_target.number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'

        c_var = ws_records.cell(row=current_row, column=6, value=f"=D{current_row}-E{current_row}")
        c_var.number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'

        c_pct = ws_records.cell(row=current_row, column=7, value=f"=IF(E{current_row}>0, D{current_row}/E{current_row}, 0)")
        c_pct.number_format = '0.0%'

        c_txns = ws_records.cell(row=current_row, column=8, value=txns)
        c_txns.number_format = '#,##0'

        c_atv = ws_records.cell(row=current_row, column=9, value=f"=IF(H{current_row}>0, D{current_row}/H{current_row}, 0)")
        c_atv.number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'

        ws_records.cell(row=current_row, column=10, value=r.get("footfall", 0)).number_format = '#,##0'
        ws_records.cell(row=current_row, column=11, value=r.get("submitted_by_name", ""))
        ws_records.cell(row=current_row, column=12, value=r.get("notes", ""))

        for c in range(1, len(headers) + 1):
            ws_records.cell(row=current_row, column=c).border = thin_border
            ws_records.cell(row=current_row, column=c).font = normal_font

        current_row += 1

    # Totals row if there are records
    if records:
        total_row = current_row
        ws_records.cell(row=total_row, column=1, value="TOTAL").font = bold_font
        ws_records.cell(row=total_row, column=1).alignment = Alignment(horizontal="center")

        t_actual = ws_records.cell(row=total_row, column=4, value=f"=SUM(D5:D{total_row-1})")
        t_actual.font = bold_font
        t_actual.number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'

        t_target = ws_records.cell(row=total_row, column=5, value=f"=SUM(E5:E{total_row-1})")
        t_target.font = bold_font
        t_target.number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'

        t_var = ws_records.cell(row=total_row, column=6, value=f"=D{total_row}-E{total_row}")
        t_var.font = bold_font
        t_var.number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'

        t_pct = ws_records.cell(row=total_row, column=7, value=f"=IF(E{total_row}>0, D{total_row}/E{total_row}, 0)")
        t_pct.font = bold_font
        t_pct.number_format = '0.0%'

        t_txns = ws_records.cell(row=total_row, column=8, value=f"=SUM(H5:H{total_row-1})")
        t_txns.font = bold_font
        t_txns.number_format = '#,##0'

        t_atv = ws_records.cell(row=total_row, column=9, value=f"=IF(H{total_row}>0, D{total_row}/H{total_row}, 0)")
        t_atv.font = bold_font
        t_atv.number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'

        for c in range(1, len(headers) + 1):
            cell = ws_records.cell(row=total_row, column=c)
            cell.border = total_border

    # Adjust column widths
    for col in ws_records.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and "\n" not in val_str:
                max_len = len(val_str)
        ws_records.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # 2. Sheet 2: Store Master & Targets
    ws_stores = wb.create_sheet(title="Store Master")
    ws_stores.views.sheetView[0].showGridLines = True
    ws_stores["A1"] = "CONCEPT STORES CONFIGURATION"
    ws_stores["A1"].font = title_font

    store_headers = ["Store Code", "Store Name", "Manager Username", f"Default Daily Target ({CURRENCY_SYMBOL})", "Status"]
    for idx, sh in enumerate(store_headers, start=1):
        c = ws_stores.cell(row=3, column=idx, value=sh)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    stores = db.get_all_stores(active_only=False)
    for idx, st in enumerate(stores, start=4):
        mgr = (st.get("manager_username") or "").strip()
        mgr_val = f"@{mgr.lstrip('@')}" if mgr else "Unassigned"
        ws_stores.cell(row=idx, column=1, value=st["code"]).alignment = Alignment(horizontal="center")
        ws_stores.cell(row=idx, column=2, value=st["name"])
        ws_stores.cell(row=idx, column=3, value=mgr_val).alignment = Alignment(horizontal="center")
        t_cell = ws_stores.cell(row=idx, column=4, value=st["default_target"])
        t_cell.number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'
        ws_stores.cell(row=idx, column=5, value="Active" if st["is_active"] else "Inactive").alignment = Alignment(horizontal="center")

    for col in ws_stores.columns:
        col_letter = get_column_letter(col[0].column)
        ws_stores.column_dimensions[col_letter].width = 22

    # Save file
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_path = EXPORTS_DIR / f"Concept_Stores_Sales_Report_{timestamp}.xlsx"
    wb.save(str(file_path))
    return file_path
